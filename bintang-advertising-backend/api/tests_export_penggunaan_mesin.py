"""Ekspor Excel Penggunaan Mesin: milik sendiri (staff) dan semua (owner/manager)."""
import io
from datetime import timedelta

import openpyxl
from django.contrib.auth import get_user_model
from django.utils import timezone
from rest_framework.test import APITestCase

from .models import Mesin, PenggunaanMesin

URL_SAYA = '/api/penggunaan-mesin/export-saya/'
URL_SEMUA = '/api/penggunaan-mesin/export/'


def _baca(response):
    assert response.status_code == 200, response.content
    return openpyxl.load_workbook(io.BytesIO(response.content))


class DasarExportMesin(APITestCase):
    def setUp(self):
        U = get_user_model()
        self.staff = U.objects.create_user(username='staff_mesin', password='x', role='staff')
        self.staff_lain = U.objects.create_user(username='staff_lain', password='x', role='staff')
        self.owner = U.objects.create_user(username='owner_mesin', password='x', role='owner')
        self.a3 = Mesin.objects.create(nama='A3 Uji', tipe='docucolor')
        self.banner = Mesin.objects.create(nama='Banner Uji', tipe='cetak_banner', basis_pencatatan='meter')
        sekarang = timezone.now()

        def catat(operator, mesin, **kw):
            return PenggunaanMesin.objects.create(mesin=mesin, operator=operator, waktu=kw.pop('waktu', sekarang), **kw)

        self.k1 = catat(self.staff, self.a3, lembar_color=10, lembar_mono=5,
                        ukuran_kertas='A3', jenis_kertas='Art Paper', gramasi_kertas='120gr')
        self.k2 = catat(self.staff, self.a3, lembar_color=3, lembar_mono=2, catatan_konfirmasi='=HYPERLINK("http://x")')
        self.k3 = catat(self.staff, self.banner, panjang_bahan_meter='2.50', jenis_bahan='Flexi 440',
                        kondisi_hasil='kendala', waktu=sekarang - timedelta(days=10))
        self.milik_lain = catat(self.staff_lain, self.a3, lembar_color=99, lembar_mono=99)


class ExportSayaTests(DasarExportMesin):
    def test_wajib_login(self):
        self.assertEqual(self.client.get(URL_SAYA).status_code, 401)

    def test_staff_hanya_mengekspor_catatan_sendiri(self):
        self.client.force_authenticate(self.staff)
        wb = _baca(self.client.get(URL_SAYA))
        ws = wb['Log Penggunaan Mesin']
        staff_kolom = {ws.cell(r, 3).value for r in range(2, ws.max_row)}  # tanpa baris TOTAL
        self.assertEqual(staff_kolom, {'staff_mesin'})
        self.assertEqual(ws.max_row - 2, 3)  # 3 entri + header, sisanya baris TOTAL

    def test_meminta_operator_orang_lain_hasilnya_kosong_tidak_bocor(self):
        self.client.force_authenticate(self.staff)
        wb = _baca(self.client.get(URL_SAYA, {'operator': self.staff_lain.id}))
        ws = wb['Log Penggunaan Mesin']
        self.assertEqual(ws.max_row, 2)   # hanya header + baris TOTAL, tanpa data
        semua_teks = ' '.join(str(c.value) for row in ws.iter_rows() for c in row)
        self.assertNotIn('staff_lain', semua_teks)
        self.assertNotIn('99', [str(ws.cell(r, 5).value) for r in range(2, ws.max_row + 1)])
        self.assertEqual(wb['Ringkasan per Mesin'].max_row, 1)  # ringkasan pun kosong

    def test_meminta_operator_diri_sendiri_tetap_berjalan(self):
        # Panel staff di frontend selalu mengirim operator=id sendiri.
        self.client.force_authenticate(self.staff)
        ws = _baca(self.client.get(URL_SAYA, {'operator': self.staff.id}))['Log Penggunaan Mesin']
        self.assertEqual(ws.max_row - 2, 3)

    def test_kolom_kertas_total_dan_ringkasan_per_mesin(self):
        self.client.force_authenticate(self.staff)
        wb = _baca(self.client.get(URL_SAYA))
        ws = wb['Log Penggunaan Mesin']
        header = [c.value for c in ws[1]]
        for kolom in ('Ukuran Kertas', 'Jenis Kertas', 'Gramasi Kertas'):
            self.assertIn(kolom, header)
        baris = {ws.cell(r, 1).value: [c.value for c in ws[r]] for r in range(1, ws.max_row + 1)}
        total = baris['TOTAL']
        self.assertEqual((total[4], total[5]), (13, 7))          # 10+3 color, 5+2 mono
        self.assertEqual(total[9], 2.5)                           # meter
        ringkas = wb['Ringkasan per Mesin']
        data = {ringkas.cell(r, 1).value: [c.value for c in ringkas[r]] for r in range(2, ringkas.max_row + 1)}
        self.assertEqual(data['A3 Uji'][1:], [2, 13, 7, 20, 0])   # entri, color, mono, klik, meter
        self.assertEqual(data['Banner Uji'][1], 1)
        self.assertEqual(data['Banner Uji'][5], 2.5)

    def test_filter_mesin_dan_tanggal_berlaku(self):
        self.client.force_authenticate(self.staff)
        ws = _baca(self.client.get(URL_SAYA, {'mesin': self.banner.id}))['Log Penggunaan Mesin']
        self.assertEqual(ws.max_row - 2, 1)
        mulai = (timezone.localdate() - timedelta(days=3)).isoformat()
        ws = _baca(self.client.get(URL_SAYA, {'tanggal_mulai': mulai}))['Log Penggunaan Mesin']
        self.assertEqual(ws.max_row - 2, 2)   # entri 10 hari lalu tersaring

    def test_teks_berawalan_sama_dengan_tidak_menjadi_rumus(self):
        self.client.force_authenticate(self.staff)
        ws = _baca(self.client.get(URL_SAYA))['Log Penggunaan Mesin']
        sel = next(c for row in ws.iter_rows(min_row=2) for c in row if isinstance(c.value, str) and 'HYPERLINK' in c.value)
        self.assertEqual(sel.data_type, 's')   # teks, bukan 'f' (rumus)

    def test_nama_file_dan_tipe(self):
        self.client.force_authenticate(self.staff)
        r = self.client.get(URL_SAYA)
        self.assertIn('penggunaan-mesin-saya-', r['Content-Disposition'])
        self.assertIn('spreadsheetml', r['Content-Type'])


class ExportSemuaTetapBerjalanTests(DasarExportMesin):
    """Ekspor Owner/Manager yang sudah ada: tetap semua data, tetap khusus owner/manager."""

    def test_owner_melihat_semua_staff_dan_kolom_baru(self):
        self.client.force_authenticate(self.owner)
        ws = _baca(self.client.get(URL_SEMUA))['Log Penggunaan Mesin']
        self.assertEqual(ws.max_row - 2, 4)
        self.assertEqual({ws.cell(r, 3).value for r in range(2, ws.max_row)}, {'staff_mesin', 'staff_lain'})
        self.assertIn('Jenis Kertas', [c.value for c in ws[1]])

    def test_staff_tidak_boleh_memakai_ekspor_semua(self):
        self.client.force_authenticate(self.staff)
        self.assertEqual(self.client.get(URL_SEMUA).status_code, 403)
