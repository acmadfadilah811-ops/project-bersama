"""Penggunaan Mesin -- Master Mesin, Log Penggunaan, Riwayat Maintenance
(lihat api/machine_models.py untuk konteks lengkap fitur ini)."""
import openpyxl
from django.db.models import Count, Sum
from django.http import HttpResponse
from django.utils import timezone
from rest_framework import viewsets
from rest_framework.decorators import action
from rest_framework.permissions import IsAuthenticated, BasePermission, SAFE_METHODS
from rest_framework.response import Response

from ..models import Mesin, PenggunaanMesin, MaintenanceMesin
from ..serializers import (
    MesinSerializer, PenggunaanMesinSerializer, MaintenanceMesinSerializer,
)
from ..permissions import IsOwnerManagerAdminOrReadOnly, IsOwnerOrManager


class IsOwnerManagerAdminOrOwnEntryReadCreate(BasePermission):
    """Log Penggunaan Mesin: siapa pun yang login boleh mencatat (staff di
    lini produksi adalah yang paling sering mengisi ini), tapi ubah/hapus
    entri hanya boleh oleh Owner/Manager/Admin atau operator yang membuat
    entri itu sendiri (mis. salah input, langsung dikoreksi)."""

    def has_permission(self, request, view):
        return bool(request.user and request.user.is_authenticated)

    def has_object_permission(self, request, view, obj):
        if request.method in SAFE_METHODS:
            return True
        if getattr(request.user, 'role', '') in ['owner', 'manager', 'admin']:
            return True
        return obj.operator_id == request.user.id


class MesinViewSet(viewsets.ModelViewSet):
    queryset = Mesin.objects.select_related('divisi').all()
    serializer_class = MesinSerializer
    permission_classes = [IsAuthenticated, IsOwnerManagerAdminOrReadOnly]

    def get_queryset(self):
        queryset = super().get_queryset()
        tipe = self.request.query_params.get('tipe')
        if tipe:
            queryset = queryset.filter(tipe=tipe)
        aktif = self.request.query_params.get('is_active')
        if aktif is not None:
            queryset = queryset.filter(is_active=aktif in ('true', '1', 'True'))
        return queryset


class PenggunaanMesinViewSet(viewsets.ModelViewSet):
    queryset = PenggunaanMesin.objects.select_related('mesin', 'operator', 'job').all()
    serializer_class = PenggunaanMesinSerializer
    permission_classes = [IsOwnerManagerAdminOrOwnEntryReadCreate]

    def get_queryset(self):
        queryset = super().get_queryset()
        mesin_id = self.request.query_params.get('mesin')
        if mesin_id:
            queryset = queryset.filter(mesin_id=mesin_id)
        job_id = self.request.query_params.get('job')
        if job_id:
            queryset = queryset.filter(job_id=job_id)
        # ?operator= -- filter per staff, dasar tab "Log Penggunaan Mesin"
        # Owner/Manager (bisa lihat & pertanggungjawabkan per orang) dan juga
        # dipakai laporan riwayat pekerjaan staff sendiri (fitur 2026-09-09).
        operator_id = self.request.query_params.get('operator')
        if operator_id:
            queryset = queryset.filter(operator_id=operator_id)
        tanggal_mulai = self.request.query_params.get('tanggal_mulai')
        tanggal_akhir = self.request.query_params.get('tanggal_akhir')
        if tanggal_mulai:
            queryset = queryset.filter(waktu__date__gte=tanggal_mulai)
        if tanggal_akhir:
            queryset = queryset.filter(waktu__date__lte=tanggal_akhir)
        return queryset

    @action(detail=False, methods=['get'], url_path='ringkasan-staff', permission_classes=[IsOwnerOrManager])
    def ringkasan_staff(self, request):
        """GET /penggunaan-mesin/ringkasan-staff/?mesin=&tanggal_mulai=&tanggal_akhir=

        Total pemakaian per staff (operator) dalam rentang tanggal -- dasar
        panel "Ringkasan per Staff" di Penggunaan Mesin (Owner/Manager),
        supaya totalnya akurat dari SELURUH data yang match filter, bukan
        cuma dijumlah dari 1 halaman tabel yang mungkin terpotong paginasi.
        """
        queryset = self.get_queryset().exclude(operator__isnull=True)
        rows = (
            queryset
            .values('operator_id', 'operator__username', 'operator__first_name', 'operator__last_name')
            .annotate(
                total_lembar_color=Sum('lembar_color'),
                total_lembar_mono=Sum('lembar_mono'),
                total_meter=Sum('panjang_bahan_meter'),
                jumlah_entri=Count('id'),
            )
            .order_by('operator__username')
        )
        result = [
            {
                'operator_id': row['operator_id'],
                'operator_nama': (
                    f"{row['operator__first_name']} {row['operator__last_name']}".strip()
                    or row['operator__username']
                ),
                'total_klik': (row['total_lembar_color'] or 0) + (row['total_lembar_mono'] or 0),
                'total_lembar_color': row['total_lembar_color'] or 0,
                'total_lembar_mono': row['total_lembar_mono'] or 0,
                'total_meter': float(row['total_meter'] or 0),
                'jumlah_entri': row['jumlah_entri'],
            }
            for row in rows
        ]
        return Response(result)

    @action(detail=False, methods=['get'], url_path='export', permission_classes=[IsOwnerOrManager])
    def export(self, request):
        """GET /penggunaan-mesin/export/?mesin=&operator=&tanggal_mulai=&tanggal_akhir=

        Excel Log Penggunaan -- pakai filter yang sama persis dengan tabel
        Log Penggunaan di layar (get_queryset()), supaya hasil export selalu
        cocok dengan yang sedang ditampilkan (fitur 2026-09-09).
        """
        return _respons_excel(self.get_queryset().order_by('-waktu'), 'log-penggunaan-mesin')

    @action(detail=False, methods=['get'], url_path='export-saya', permission_classes=[IsAuthenticated])
    def export_saya(self, request):
        """GET /penggunaan-mesin/export-saya/?mesin=&tanggal_mulai=&tanggal_akhir=

        Excel penggunaan mesin MILIK PEMINTA SENDIRI -- bahan laporan staff.
        Operator DIKUNCI ke user yang login di server. Bila klien mengirim
        `operator` orang lain, filternya bertemu kunci itu dan hasilnya file
        KOSONG -- tidak pernah memuat catatan orang lain (API1: pembatasan di
        tampilan saja bukan keamanan)."""
        queryset = self.get_queryset().filter(operator=request.user).order_by('-waktu')
        return _respons_excel(queryset, 'penggunaan-mesin-saya')


HEADER_EXCEL = [
    'Tanggal', 'Mesin', 'Staff', 'Job', 'Lembar Color', 'Lembar Mono',
    'Ukuran Kertas', 'Jenis Kertas', 'Gramasi Kertas',
    'Panjang Bahan (m)', 'Jenis Bahan', 'Kondisi Hasil', 'Catatan Konfirmasi',
]
HEADER_RINGKASAN = ['Mesin', 'Jumlah Entri', 'Total Lembar Color', 'Total Lembar Mono', 'Total Klik', 'Total Panjang Bahan (m)']


def _tambah_baris(ws, nilai):
    """Tulis satu baris. Teks bebas dari pengguna (catatan, jenis bahan) yang diawali
    '=' dipaksa jadi TEKS -- openpyxl otherwise menyimpannya sebagai rumus Excel
    (rumus dari data yang tidak tepercaya)."""
    ws.append(nilai)
    for sel in ws[ws.max_row]:
        if isinstance(sel.value, str) and sel.value.startswith('='):
            sel.data_type = 's'


def _respons_excel(queryset, nama_dasar):
    """Excel penggunaan mesin: sheet detail (+ baris TOTAL) dan sheet ringkasan per mesin."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Log Penggunaan Mesin'
    _tambah_baris(ws, HEADER_EXCEL)

    total_color = total_mono = 0
    total_meter = 0.0
    per_mesin = {}
    for entry in queryset:
        operator_nama = ''
        if entry.operator:
            operator_nama = entry.operator.get_full_name() or entry.operator.username
        meter = float(entry.panjang_bahan_meter) if entry.panjang_bahan_meter is not None else None
        _tambah_baris(ws, [
            timezone.localtime(entry.waktu).strftime('%Y-%m-%d %H:%M'),
            entry.mesin.nama if entry.mesin else '-',
            operator_nama,
            entry.job.nomor_sumber if entry.job else '-',
            entry.lembar_color,
            entry.lembar_mono,
            entry.ukuran_kertas,
            entry.jenis_kertas,
            entry.gramasi_kertas,
            meter if meter is not None else '',
            entry.jenis_bahan,
            'Ada Kendala' if entry.kondisi_hasil == 'kendala' else 'OK',
            entry.catatan_konfirmasi,
        ])
        total_color += entry.lembar_color
        total_mono += entry.lembar_mono
        total_meter += meter or 0
        r = per_mesin.setdefault(entry.mesin.nama if entry.mesin else '-', [0, 0, 0, 0.0])
        r[0] += 1
        r[1] += entry.lembar_color
        r[2] += entry.lembar_mono
        r[3] += meter or 0
    _tambah_baris(ws, ['TOTAL', '', '', '', total_color, total_mono, '', '', '', round(total_meter, 2), '', '', ''])

    ringkas = wb.create_sheet('Ringkasan per Mesin')
    _tambah_baris(ringkas, HEADER_RINGKASAN)
    for nama_mesin, (n, color, mono, meter) in sorted(per_mesin.items()):
        _tambah_baris(ringkas, [nama_mesin, n, color, mono, color + mono, round(meter, 2)])

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{nama_dasar}-{timezone.localdate():%Y%m%d}.xlsx"'
    wb.save(response)
    return response


class MaintenanceMesinViewSet(viewsets.ModelViewSet):
    queryset = MaintenanceMesin.objects.select_related('mesin', 'dicatat_oleh').all()
    serializer_class = MaintenanceMesinSerializer
    permission_classes = [IsAuthenticated, IsOwnerManagerAdminOrReadOnly]

    def get_queryset(self):
        queryset = super().get_queryset()
        mesin_id = self.request.query_params.get('mesin')
        if mesin_id:
            queryset = queryset.filter(mesin_id=mesin_id)
        return queryset
